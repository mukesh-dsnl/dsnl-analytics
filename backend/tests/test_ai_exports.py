"""Slash commands and /csv /excel exports: the file holds the *full* result of
the answer's queries — past the cap the model reads — with readable headers,
and is only downloadable by the conversation's owner."""

import csv
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401 — register every table
from app.ai import commands, exports
from app.ai.providers.base import LLMClient, LLMTurn, ToolCallRequest
from app.ai.tools import ad_hoc_sql, metrics
from app.core.config import get_settings
from app.core.database import Base, get_db
from app.main import app
from app.models.conversation import (
    EXPORT_FAILED,
    EXPORT_READY,
    Conversation,
    Message,
    MessageExport,
)
from app.models.user import User

# ── Commands ───────────────────────────────────────────────────────────────


def test_commands_are_parsed_out_of_the_question():
    parsed = commands.parse("/Voicedrop /excel top accounts last week /excel")
    assert parsed.commands == ["voicedrop", "excel"]
    assert parsed.text == "top accounts last week"
    assert parsed.scopes == ["voicedrop"] and parsed.exports == ["excel"]


def test_unknown_commands_and_paths_stay_as_text():
    parsed = commands.parse("show calls /foo for 2026/10/01 and/or /chart")
    assert parsed.commands == ["chart"]
    assert parsed.text == "show calls /foo for 2026/10/01 and/or"


def test_history_keeps_commands_as_a_short_tag():
    assert commands.strip("/multicall /csv minutes yesterday") == "[/multicall, /csv] minutes yesterday"
    assert commands.strip("no commands here") == "no commands here"


def test_instructions_tell_the_model_what_each_command_means():
    text = commands.instructions(commands.parse("/voicedrop /excel /chart q"))
    assert 'service="voicedrop"' in text
    assert "attached to this answer as an Excel workbook" in text
    assert "draws a chart" in text
    assert commands.instructions(commands.parse("plain question")) == ""
    two = commands.instructions(commands.parse("/voicedrop /conference q"))
    assert "only these services: Voicedrop, Conference" in two


# ── Headers ────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("column", "expected"),
    [
        ("total_minutes", "Total Minutes"),
        ("ACCOUNTID", "Account ID"),
        ("TEL_DIGIT", "Phone (Dial Out)"),
        ("START_DATETIME_EPOC", "Start Time (Epoch)"),
        ("connect_rate_pct", "Connect Rate %"),
        ("crn", "CRN"),
    ],
)
def test_headers_are_readable(column, expected):
    assert exports.header(column) == expected


# ── Building files from the lake ───────────────────────────────────────────


@pytest.fixture
def export_env(fixture_lake, tmp_path, monkeypatch):
    """The fixture lake, a tiny model cap, and a private export directory."""
    settings = get_settings()
    monkeypatch.setattr(settings, "AI_MAX_ROWS_TO_MODEL", 2)
    monkeypatch.setattr(metrics, "MAX_ROWS", 2)
    monkeypatch.setattr(settings, "AI_EXPORT_DIR", str(tmp_path / "exports"))

    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(bind=engine)
    db = sessionmaker(bind=engine)()
    yield fixture_lake, db
    db.close()
    engine.dispose()


def _message(db, queries, question="calls by phone"):
    conversation = Conversation(id="c0ffee00-0000-0000-0000-000000000001", user_id="u1", title=question)
    db.add(conversation)
    db.flush()
    message = Message(conversation_id=conversation.id, status="pass", query=question, response="ok", queries=queries)
    db.add(message)
    db.commit()
    return message


def _sql_call(day, sql="SELECT CRN, TEL_DIGIT, ACCOUNTID, START_DATETIME FROM cdr ORDER BY TEL_DIGIT"):
    return {"tool": "run_cdr_query", "error": False,
            "input": {"date_from": str(day), "date_to": str(day), "sql": sql, "purpose": "every leg"}}


def test_the_file_holds_every_row_not_just_what_the_model_saw(export_env):
    day, db = export_env
    call = _sql_call(day)

    # What the model was shown: capped at 2 and marked truncated.
    content, is_error = ad_hoc_sql.run_cdr_query(**call["input"])
    shown = json.loads(content)
    assert not is_error and shown["row_count"] == 2 and shown["truncated"]

    message = _message(db, [call])
    pending = exports.create_pending(db, message, ["xlsx", "csv"], "u1")
    db.commit()
    exports.build(db, message, pending, "calls by phone")

    xlsx, csv_export = pending
    assert xlsx.status == csv_export.status == EXPORT_READY
    assert xlsx.row_count == csv_export.row_count == 6  # every leg in the lake

    with exports.resolve(csv_export).open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    assert rows[0] == ["CRN", "Phone (Dial Out)", "Account ID", "Start Time"]
    assert len(rows) == 7

    book = load_workbook(exports.resolve(xlsx), read_only=True)
    # Only the data that was asked for — no notes sheet, nothing else.
    assert book.sheetnames == ["Data"]
    sheet = book.worksheets[0]
    values = list(sheet.iter_rows(values_only=True))
    assert values[0] == ("CRN", "Phone (Dial Out)", "Account ID", "Start Time")
    assert len(values) == 7
    # Identifiers are text, so Excel keeps every digit.
    assert all(isinstance(row[1], str) for row in values[1:])
    assert all(isinstance(row[0], str) for row in values[1:])


def test_metrics_exports_ignore_the_model_cap_but_keep_numbers_numeric(export_env):
    day, db = export_env
    call = {"tool": "query_metrics", "error": False,
            "input": {"date_from": str(day), "date_to": str(day),
                      "measures": ["calls", "minutes"], "group_by": ["conference"]}}

    shown = json.loads(metrics.query_metrics(**call["input"])[0])
    assert shown["row_count"] <= 2

    message = _message(db, [call])
    (xlsx,) = exports.create_pending(db, message, ["xlsx"], "u1")
    db.commit()
    exports.build(db, message, [xlsx], "calls by conference")

    sheet_rows = list(load_workbook(exports.resolve(xlsx), read_only=True).worksheets[0].iter_rows(values_only=True))
    assert sheet_rows[0] == ("Conference", "Calls", "Minutes")
    assert all(isinstance(row[1], int) for row in sheet_rows[1:])
    assert xlsx.row_count == len(sheet_rows) - 1


def test_both_files_hold_only_the_final_query(export_env):
    """Exploratory calls the model made on the way are not the data asked for."""
    day, db = export_env
    exploring = _sql_call(day, "SELECT CRN FROM cdr")
    final = {"tool": "query_metrics", "error": False,
             "input": {"date_from": str(day), "date_to": str(day), "measures": ["calls"]}}
    failed = {"tool": "query_metrics", "error": True, "input": {}}

    message = _message(db, [exploring, final, failed])
    xlsx, csv_export = exports.create_pending(db, message, ["xlsx", "csv"], "u1")
    db.commit()
    exports.build(db, message, [xlsx, csv_export], "total calls")

    book = load_workbook(exports.resolve(xlsx), read_only=True)
    assert book.sheetnames == ["Data"] and xlsx.sheet_count == 1
    assert next(book["Data"].iter_rows(values_only=True)) == ("Calls",)
    with exports.resolve(csv_export).open(encoding="utf-8-sig") as handle:
        assert next(csv.reader(handle)) == ["Calls"]


def test_the_previous_query_is_used_when_the_final_one_cannot_be_rerun(export_env):
    day, db = export_env
    good = _sql_call(day, "SELECT CRN FROM cdr")
    broken = _sql_call(day, "DROP TABLE cdr")
    message = _message(db, [good, broken])
    (csv_export,) = exports.create_pending(db, message, ["csv"], "u1")
    db.commit()
    exports.build(db, message, [csv_export], "crns")

    assert csv_export.status == EXPORT_READY and csv_export.row_count == 6


def test_an_answer_without_data_queries_has_nothing_to_export(export_env):
    _, db = export_env
    message = _message(db, [])
    (csv_export,) = exports.create_pending(db, message, ["csv"], "u1")
    db.commit()
    exports.build(db, message, [csv_export], "hello")

    assert csv_export.status == EXPORT_FAILED
    assert "did not query any data" in csv_export.error


def test_a_query_that_cannot_be_rerun_is_reported_not_raised(export_env):
    day, db = export_env
    broken = _sql_call(day, "DELETE FROM cdr")
    message = _message(db, [broken])
    (xlsx,) = exports.create_pending(db, message, ["xlsx"], "u1")
    db.commit()
    exports.build(db, message, [xlsx], "broken")

    assert xlsx.status == EXPORT_FAILED
    assert "could not be exported" in xlsx.error


def test_resolve_refuses_paths_outside_the_export_directory(export_env):
    _, db = export_env
    assert exports.resolve(MessageExport(file_path="../../etc/passwd")) is None


# ── Through the API ────────────────────────────────────────────────────────


class ScriptedLLM(LLMClient):
    """Calls one query on round one, answers on round two."""

    provider, model = "fake", "fake-1"

    def __init__(self, call):
        self.call = call
        self.systems: list[str] = []
        self.questions: list[str] = []

    def send(self, system, history, tools):
        self.systems.append(system)
        self.questions.append(history[-1].text if history else "")
        if len(self.systems) == 1:
            request = ToolCallRequest(id="t1", name=self.call["tool"], input=self.call["input"])
            return LLMTurn(text="", tool_calls=[request], stop_reason="tool_use")
        return LLMTurn(text="| CRN |\n|---|\n| 1 |")


@pytest.fixture
def api(export_env, monkeypatch):
    day, _ = export_env
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(bind=engine)
    factory = sessionmaker(bind=engine)

    def override_db():
        with factory() as db:
            yield db

    app.dependency_overrides[get_db] = override_db
    with factory() as db:
        db.add_all([
            User(username="owner", password_hash="pw", user_id="owner-id", ai_permission=True),
            User(username="other", password_hash="pw", user_id="other-id", ai_permission=True),
        ])
        db.commit()

    llm = ScriptedLLM(_sql_call(day))
    monkeypatch.setattr("app.api.ai.get_llm_client", lambda: llm)

    def login(name):
        client = TestClient(app)
        assert client.post("/api/auth/login", json={"username": name, "password": "pw"}).status_code == 200
        return client

    yield login, llm
    app.dependency_overrides.pop(get_db, None)
    engine.dispose()


def test_chat_with_commands_returns_a_downloadable_file(api):
    login, llm = api
    owner = login("owner")

    response = owner.post("/api/ai/chat", json={"question": "/voicedrop /csv calls by phone"})
    assert response.status_code == 200, response.text
    body = response.json()

    # The model was asked the question without the tokens, with the scope told to it.
    assert llm.questions[0] == "calls by phone"
    assert 'service="voicedrop"' in llm.systems[0]

    (export,) = body["exports"]
    assert export["status"] == "ready" and export["format"] == "csv" and export["row_count"] == 6
    download = owner.get(export["download_url"])
    assert download.status_code == 200
    assert download.headers["content-type"].startswith("text/csv")
    assert download.text.lstrip("﻿").splitlines()[0] == "CRN,Phone (Dial Out),Account ID,Start Time"
    assert export["file_name"].startswith("calls-by-phone-") and export["file_name"].endswith(".csv")

    # A reload shows the question as typed, and its file.
    detail = owner.get(f"/api/ai/conversations/{body['conversation_id']}").json()
    stored = detail["interactions"][0]
    assert stored["query"] == "/voicedrop /csv calls by phone"
    assert stored["exports"][0]["id"] == export["id"]


def test_someone_elses_export_is_a_404(api):
    login, _ = api
    body = login("owner").post("/api/ai/chat", json={"question": "/excel calls"}).json()
    url = body["exports"][0]["download_url"]

    assert login("other").get(url).status_code == 404
    assert login("owner").get("/api/ai/exports/999999/download").status_code == 404


# ── Preview limits are not data limits ─────────────────────────────────────


def test_a_preview_limit_the_user_did_not_ask_for_is_removed(export_env):
    """The reported bug: asked for /excel, the model wrote LIMIT 20 to keep its
    preview short, and the file held 20 rows instead of every call."""
    day, db = export_env
    call = _sql_call(day, "SELECT CRN, TEL_DIGIT FROM cdr ORDER BY TEL_DIGIT LIMIT 2")
    message = _message(db, [call], question="give me the connected list")
    (xlsx,) = exports.create_pending(db, message, ["xlsx"], "u1")
    db.commit()
    exports.build(db, message, [xlsx], "give me the connected list")

    assert xlsx.row_count == 6
    assert load_workbook(exports.resolve(xlsx), read_only=True).sheetnames == ["Data"]


def test_a_limit_the_question_names_is_kept(export_env):
    day, db = export_env
    call = _sql_call(day, "SELECT CRN FROM cdr ORDER BY CRN LIMIT 3")
    message = _message(db, [call], question="top 3 conferences")
    (csv_export,) = exports.create_pending(db, message, ["csv"], "u1")
    db.commit()
    exports.build(db, message, [csv_export], "top 3 conferences")
    assert csv_export.row_count == 3


@pytest.mark.parametrize(
    ("sql", "question", "kept"),
    [
        # The exact statement from the reported conversation (message 186).
        ("SELECT c.CALL_DATE AS Date, c.CRN FROM cdr c JOIN codr o ON o.CRN = c.CRN "
         "WHERE c.INCONF_DATETIME_EPOC <> 0 ORDER BY c.START_DATETIME_EPOC LIMIT 20",
         "give me the oct 1 connected list of their Date, CRN, Phone Number, Duration", False),
        ("SELECT ACCOUNTID, COUNT(*) FROM cdr GROUP BY 1 ORDER BY 2 DESC LIMIT 10", "top 10 accounts", True),
        ("SELECT * FROM (SELECT CRN FROM cdr LIMIT 5) t", "list crns", True),  # inner limit untouched
        ("SELECT CRN FROM cdr LIMIT 20;", "oct 1 2026 list", False),  # "1" and "2026" are not 20
        # A day of the month is not a row count.
        ("SELECT CRN FROM cdr LIMIT 10", "oct 10 connected list", False),
        ("SELECT CRN FROM cdr LIMIT 10", "calls on 10th and 12th", False),
        ("SELECT CRN FROM cdr LIMIT 3", "oct 1 to 3 calls", False),
        ("SELECT CRN FROM cdr LIMIT 5", "first 5 calls on 2026-10-05", True),
    ],
)
def test_which_limits_count_as_preview_limits(sql, question, kept):
    arguments, note = exports.full_arguments({"tool": "run_cdr_query", "input": {"sql": sql}}, question)
    assert (arguments["sql"] == sql) is kept
    assert (note is None) is kept


def test_a_metrics_limit_is_treated_the_same_way():
    call = {"tool": "query_metrics", "input": {"measures": ["calls"], "limit": 20}}
    assert "limit" not in exports.full_arguments(call, "calls by account")[0]
    assert exports.full_arguments(call, "top 20 accounts by calls")[0]["limit"] == 20
